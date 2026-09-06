"""
简历文件解析器测试
覆盖 md/txt 解码、docx 段落表格提取、手工排版章节识别、pdf 复用 PDFParser、非法输入报错
"""
import io

import pytest
from unittest.mock import MagicMock

from src.knowledge.resume_file_parser import (
    ResumeFileParser, SUPPORTED_EXTENSIONS, is_bold_section_title,
)


def _build_docx_bytes(blocks):
    """内存构造docx：blocks为("heading"/"para"/"list"/"table"/"bold_para"/"mixed_para", 内容)列表"""
    from docx import Document

    doc = Document()
    for kind, content in blocks:
        if kind == "heading":
            doc.add_heading(content, level=1)
        elif kind == "para":
            doc.add_paragraph(content)
        elif kind == "bold_para":
            # 手工排版章节标题：普通段落整段加粗
            doc.add_paragraph().add_run(content).bold = True
        elif kind == "mixed_para":
            # 部分加粗的正文行（不应识别为章节标题）
            para = doc.add_paragraph()
            para.add_run(content).bold = True
            para.add_run("普通尾巴")
        elif kind == "list":
            doc.add_paragraph(content, style="List Bullet")
        elif kind == "table":
            table = doc.add_table(rows=len(content), cols=len(content[0]))
            for i, row in enumerate(content):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestParseText:
    """md/txt直接解码"""

    def test_parse_md_utf8(self):
        parser = ResumeFileParser()
        result = parser.parse("resume.md", "# 张三\n\n前端工程师".encode("utf-8"))
        assert "# 张三" in result
        assert "前端工程师" in result

    def test_parse_txt_gbk_fallback(self):
        """utf-8解码失败时应退回gbk（Windows中文txt常见编码）"""
        parser = ResumeFileParser()
        content = "个人信息：张三，5年前端经验"
        result = parser.parse("resume.txt", content.encode("gbk"))
        assert "张三" in result
        assert "5年前端经验" in result

    def test_empty_text_raises(self):
        parser = ResumeFileParser()
        with pytest.raises(ValueError, match="解析结果为空"):
            parser.parse("resume.md", b"   \n  ")


class TestUnsupportedType:
    """非法文件类型"""

    def test_html_raises(self):
        parser = ResumeFileParser()
        with pytest.raises(ValueError, match="不支持的简历文件格式"):
            parser.parse("resume.html", b"<html></html>")

    def test_no_extension_raises(self):
        parser = ResumeFileParser()
        with pytest.raises(ValueError, match="不支持的简历文件格式"):
            parser.parse("resume", b"plain text")

    def test_supported_extensions_covers_four_types(self):
        assert set(SUPPORTED_EXTENSIONS) == {".md", ".txt", ".docx", ".pdf"}


class TestParseDocx:
    """docx段落与表格提取"""

    def test_parse_heading_paragraph_list(self):
        parser = ResumeFileParser()
        data = _build_docx_bytes([
            ("heading", "张三的简历"),
            ("para", "5年前端开发经验，正在转向AI Agent方向"),
            ("list", "主导简历生成Agent项目"),
        ])
        result = parser.parse("resume.docx", data)
        assert "# 张三的简历" in result          # Heading 1 -> "# "
        assert "5年前端开发经验" in result
        assert "- 主导简历生成Agent项目" in result  # List Bullet -> "- "

    def test_parse_table_rows(self):
        parser = ResumeFileParser()
        data = _build_docx_bytes([
            ("table", [["公司", "职位"], ["某科技", "前端工程师"]]),
        ])
        result = parser.parse("resume.docx", data)
        assert "| 公司 | 职位 |" in result
        assert "| 某科技 | 前端工程师 |" in result

    def test_parse_keeps_block_order(self):
        """段落与表格应保持文档原始顺序"""
        parser = ResumeFileParser()
        data = _build_docx_bytes([
            ("para", "第一段"),
            ("table", [["A", "B"]]),
            ("para", "第二段"),
        ])
        result = parser.parse("resume.docx", data)
        assert result.index("第一段") < result.index("| A | B |") < result.index("第二段")

    def test_empty_docx_raises(self):
        parser = ResumeFileParser()
        data = _build_docx_bytes([])
        with pytest.raises(ValueError, match="解析结果为空"):
            parser.parse("resume.docx", data)


class TestIsBoldSectionTitle:
    """手工排版章节标题判定"""

    def test_short_bold_line_is_section(self):
        assert is_bold_section_title("教育背景", True) is True

    def test_not_bold_is_not_section(self):
        assert is_bold_section_title("教育背景", False) is False

    def test_punctuated_line_is_not_section(self):
        assert is_bold_section_title("求职意向：web前端开发", True) is False

    def test_long_line_is_not_section(self):
        assert is_bold_section_title("负责公司金析云资金交易智能分析研判产品的技术调研开发测试", True) is False

    def test_first_line_relaxes_length_limit(self):
        title = "孟龙翔_web前端开发_五年经验求职简历"
        assert is_bold_section_title(title, True) is False
        assert is_bold_section_title(title, True, is_first=True) is True


class TestParseDocxManualLayout:
    """手工排版docx（章节名为加粗普通段）的标题识别"""

    def test_bold_first_line_becomes_h1(self):
        parser = ResumeFileParser()
        data = _build_docx_bytes([
            ("bold_para", "孟龙翔_web前端开发"),
            ("para", "姓名：孟龙翔"),
        ])
        result = parser.parse("resume.docx", data)
        assert "# 孟龙翔_web前端开发" in result

    def test_bold_section_lines_become_h2(self):
        parser = ResumeFileParser()
        data = _build_docx_bytes([
            ("bold_para", "孟龙翔的简历"),
            ("bold_para", "教育背景"),
            ("para", "就读院校：武汉工程大学"),
            ("bold_para", "项目经验"),
            ("para", "金析云平台"),
        ])
        result = parser.parse("resume.docx", data)
        assert "## 教育背景" in result
        assert "## 项目经验" in result
        assert "就读院校：武汉工程大学" in result

    def test_colon_and_long_bold_lines_not_sections(self):
        parser = ResumeFileParser()
        data = _build_docx_bytes([
            ("bold_para", "简历标题"),
            ("bold_para", "求职意向：web前端开发工程师"),
        ])
        result = parser.parse("resume.docx", data)
        assert "求职意向：web前端开发工程师" in result
        assert "## 求职意向" not in result

    def test_partial_bold_line_not_section(self):
        parser = ResumeFileParser()
        data = _build_docx_bytes([
            ("bold_para", "简历标题"),
            ("mixed_para", "公司"),
        ])
        result = parser.parse("resume.docx", data)
        assert "## 公司" not in result

    def test_heading_style_takes_precedence(self):
        """已用Heading样式的段落仍按样式映射，不受加粗判定影响"""
        parser = ResumeFileParser()
        data = _build_docx_bytes([
            ("heading", "张三的简历"),
            ("bold_para", "专业技能"),
        ])
        result = parser.parse("resume.docx", data)
        assert "# 张三的简历" in result
        assert "## 专业技能" in result


class TestParsePdf:
    """pdf复用PDFParser（注入mock，不触发真实MinerU）"""

    def test_parse_pdf_delegates_to_pdf_parser(self):
        mock_pdf_parser = MagicMock()
        mock_pdf_parser.parse_single.return_value = "# PDF简历\n\n张三"
        parser = ResumeFileParser(pdf_parser=mock_pdf_parser)

        result = parser.parse("resume.pdf", b"%PDF-1.4 fake")

        assert result == "# PDF简历\n\n张三"
        mock_pdf_parser.parse_single.assert_called_once()
        # 传入的是临时文件路径，保留.pdf后缀
        tmp_path = mock_pdf_parser.parse_single.call_args[0][0]
        assert tmp_path.suffix == ".pdf"
        assert not tmp_path.exists()  # 临时目录已清理

    def test_pdf_parser_error_raises_value_error(self):
        mock_pdf_parser = MagicMock()
        mock_pdf_parser.parse_single.side_effect = RuntimeError("MinerU崩了")
        parser = ResumeFileParser(pdf_parser=mock_pdf_parser)

        with pytest.raises(ValueError, match="简历PDF解析失败"):
            parser.parse("resume.pdf", b"%PDF-1.4 fake")

    def test_pdf_empty_result_raises(self):
        mock_pdf_parser = MagicMock()
        mock_pdf_parser.parse_single.return_value = "   "
        parser = ResumeFileParser(pdf_parser=mock_pdf_parser)

        with pytest.raises(ValueError, match="解析结果为空"):
            parser.parse("resume.pdf", b"%PDF-1.4 fake")
